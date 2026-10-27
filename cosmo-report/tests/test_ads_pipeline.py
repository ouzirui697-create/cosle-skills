"""Synthetic regression cases. No seller data belongs in these fixtures."""
import copy
import json
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
from parse_prompts import parse_sheet, parse_file, number
from analyze_ads import analyze
from render_ads import render

HEADERS = ['Date', 'Campaign Name', 'Ad Group Name', 'Country', 'Currency',
           'Prompt details', 'Advertised ASIN', 'Advertised SKU', 'Impressions',
           'Clicks', 'Cost Per Click (CPC)', 'Spend', '7 Day Total Orders (#)',
           '7 Day Total Sales', '7 Day Advertised SKU Orders (#)',
           '7 Day Other SKU Orders (#)', '7 Day Advertised SKU Sales', '7 Day Other SKU Sales']
ROW = ['2026-04-01', 'Campaign A', 'Group A', 'United States', 'USD',
       'Why choose Example routers?', 'B0EXAMPLE01', 'SKU-A', 1000, 40, 0.5, 20, 2, 200, 1, 1, 120, 80]


def fixture(rows=None, headers=None, name='main'):
    sheet = parse_sheet(name, [headers or HEADERS, *(rows if rows is not None else [ROW])], 'synthetic')
    return {'schema_version': 1, 'source_file': 'synthetic.xlsx', 'sha256': 'synthetic', 'sheets': [sheet]}


def paid(parsed):
    return analyze(parsed, billing='paid', target_acos=0.25, min_clicks=30, as_of='2026-04-10')


class PipelineTests(unittest.TestCase):
    def test_cpc_never_shadows_spend_and_split_is_preserved(self):
        data = paid(fixture())
        self.assertEqual(data['summary']['metrics']['spend'], 20)
        self.assertEqual(data['summary']['metrics']['cpc'], 0.5)
        self.assertEqual(data['summary']['metrics']['promoted_orders'], 1)
        self.assertEqual(data['summary']['metrics']['other_orders'], 1)
        self.assertEqual(data['units'][0]['gate']['status'], 'review_candidate')
        self.assertIsNone(data['units'][0]['gate']['bid_change_percent'])

    def test_header_permutations_keep_meaning(self):
        for offset in range(len(HEADERS)):
            h = HEADERS[offset:] + HEADERS[:offset]
            r = ROW[offset:] + ROW[:offset]
            self.assertEqual(paid(fixture([r], h))['summary']['metrics']['spend'], 20)

    def test_free_converting_l3_has_no_bid_or_pause_action(self):
        row = ROW.copy(); row[10:12] = [0, 0]
        data = analyze(fixture([row]), billing='free')
        gate = data['units'][0]['gate']
        self.assertIn('free_period', gate['reasons'])
        self.assertEqual(gate['status'], 'observe_or_verify')
        self.assertIsNone(gate['bid_change_percent'])
        self.assertIsNone(data['summary']['metrics']['roas'])
        self.assertEqual(data['summary']['metrics']['acos'], 0)

    def test_missing_spend_not_zero_or_partial_total(self):
        row = ROW.copy(); row[0] = '2026-04-02'; row[11] = None
        data = paid(fixture([ROW, row]))
        self.assertIsNone(data['summary']['metrics']['spend'])
        self.assertEqual(data['summary']['coverage']['spend']['observed_sum'], 20)
        self.assertIn('spend_incomplete', data['units'][0]['gate']['reasons'])

    def test_missing_split_does_not_invent_promoted_orders(self):
        row = ROW.copy(); row[14] = None
        data = paid(fixture([row]))
        self.assertIsNone(data['summary']['metrics']['promoted_orders'])
        self.assertIn('promoted_orders_incomplete', data['units'][0]['gate']['reasons'])

    def test_split_inconsistency_blocks_review(self):
        row = ROW.copy(); row[14] = 9
        data = paid(fixture([row]))
        self.assertIn('row_quality_issues', data['units'][0]['gate']['reasons'])

    def test_percentages_not_guessed_from_size(self):
        self.assertEqual(number('0.5%', 'acos'), 0.005)
        self.assertEqual(number(1.5, 'acos'), 1.5)
        self.assertEqual(number('150%', 'acos'), 1.5)
        self.assertEqual(number('1,234.56', 'spend'), 1234.56)
        for value in ['1.234,56', 'inf', float('nan'), -1, True]:
            with self.assertRaises(ValueError): number(value, 'spend')

    def test_campaign_or_ad_label_never_becomes_title(self):
        data = fixture(headers=HEADERS+['Ad name'], rows=[ROW+['Internal Campaign Label']])
        rec = data['sheets'][0]['records'][0]
        self.assertNotIn('title', rec)
        self.assertEqual(rec['ad_name_kind'], 'unverified_ad_label')

    def test_supplement_retained_but_not_summed(self):
        data = fixture()
        data['sheets'].append(parse_sheet('response', [HEADERS+['Sample response'],ROW+['Example Model B']], 'synthetic'))
        result = paid(data)
        self.assertEqual(result['summary']['metrics']['sales'], 200)
        self.assertEqual(len(result['response_records']), 1)
        self.assertEqual(result['not_summed_sheets'], ['response'])

    def test_multiple_main_sheets_need_explicit_choice(self):
        data = fixture()
        data['sheets'].append(parse_sheet('other', [HEADERS,ROW], 'synthetic'))
        with self.assertRaises(ValueError): analyze(data)
        self.assertEqual(analyze(data, ['main'])['summary']['record_count'], 1)
        with self.assertRaises(ValueError): analyze(data, ['main','other'])

    def test_duplicate_daily_rows_block_even_if_explicitly_merged(self):
        data = fixture()
        data['sheets'].append(parse_sheet('other', [HEADERS,ROW], 'synthetic'))
        result = analyze(data, ['main','other'], confirmed_disjoint=True, billing='paid')
        self.assertIn('duplicate_daily_entity', result['units'][0]['gate']['reasons'])

    def test_repeated_rows_are_flagged_not_silently_dropped(self):
        parsed = fixture([ROW,ROW])
        self.assertEqual(len(parsed['sheets'][0]['records']), 2)
        self.assertEqual(parsed['sheets'][0]['records'][1]['issues'][0]['code'], 'duplicate_row')
        self.assertIn('duplicate_daily_entity', paid(parsed)['warnings'])

    def test_mixed_scope_rejected(self):
        for col, value in [(3,'Canada'), (4,'CAD')]:
            row = ROW.copy(); row[col] = value
            with self.assertRaises(ValueError): paid(fixture([ROW,row]))

    def test_duplicate_attribution_headers_do_not_choose_first(self):
        parsed = fixture(headers=HEADERS+['14 Day Total Sales'],rows=[ROW+[300]])
        self.assertIsNone(parsed['sheets'][0]['records'][0]['sales'])
        self.assertTrue(parsed['sheets'][0]['issues'])
        self.assertIn('sheet_quality_issues', paid(parsed)['warnings'])

    def test_immature_and_missing_dates_block_review(self):
        row = ROW.copy(); row[0]='2026-04-09'
        self.assertIn('attribution_not_mature', paid(fixture([row]))['units'][0]['gate']['reasons'])
        row[0]=None
        self.assertIn('attribution_maturity_unknown', paid(fixture([row]))['units'][0]['gate']['reasons'])

    def test_same_text_different_ads_stays_separate(self):
        row = ROW.copy(); row[1]='Campaign B'
        data = paid(fixture([ROW,row]))
        self.assertEqual(len(data['units']),2)
        self.assertEqual(data['summary']['unique_prompt_count'],1)
        self.assertEqual(data['summary']['record_count'],2)

    def test_notes_missing_prompt_and_total_excluded_with_sources(self):
        rows=[]
        for text in ['-',None,'Total']:
            row=ROW.copy();row[5]=text;rows.append(row)
        parsed=fixture(rows)
        self.assertEqual(len(parsed['sheets'][0]['records']),0)
        self.assertEqual(len(parsed['sheets'][0]['excluded']),3)

    def test_unknown_billing_cannot_pass(self):
        self.assertIn('billing_unknown', analyze(fixture())['units'][0]['gate']['reasons'])

    def test_placeholder_identity_and_response_stay_unknown(self):
        row=ROW.copy();row[6]='-'
        data=fixture(headers=HEADERS+['Sample response'],rows=[row+['-']])
        result=analyze(data,['main'],billing='paid',target_acos=0.25,min_clicks=30,as_of='2026-04-10')
        self.assertEqual(result['response_records'],[])
        self.assertIn('action_entity_incomplete',result['units'][0]['gate']['reasons'])

    def test_small_sample_cannot_pass(self):
        row=ROW.copy();row[9]=1
        self.assertIn('sample_below_user_policy', paid(fixture([row]))['units'][0]['gate']['reasons'])

    def test_csv_bom_and_tsv(self):
        import csv
        with tempfile.TemporaryDirectory() as tmp:
            for suffix,delimiter in [('.csv',','),('.tsv','\t')]:
                path=Path(tmp)/('example'+suffix)
                with path.open('w',encoding='utf-8-sig',newline='') as f:
                    writer=csv.writer(f,delimiter=delimiter);writer.writerow(HEADERS);writer.writerow(ROW)
                self.assertEqual(parse_file(path)['sheets'][0]['records'][0]['spend'],20)

    def test_xlsx_all_sheets_even_if_active_sheet_is_notes(self):
        import openpyxl
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'synthetic.xlsx'
            wb=openpyxl.Workbook();wb.active.title='notes';wb.active.append(['Notes only'])
            ws=wb.create_sheet('data');ws.append(HEADERS);ws.append(ROW)
            response=wb.create_sheet('responses');response.append(HEADERS+['Sample response']);response.append(ROW+['Example Model B'])
            wb.save(path);wb.close()
            before=path.read_bytes();parsed=parse_file(path)
            self.assertEqual(len(parsed['sheets']),3)
            self.assertEqual(paid(parsed)['summary']['metrics']['sales'],200)
            self.assertEqual(path.read_bytes(),before)

    def test_render_escapes_input_and_enforces_annotation_identity(self):
        row=ROW.copy();row[5]='<script>alert(1)</script>'
        data=paid(fixture([row]))
        page=render(data)
        self.assertNotIn('<script>',page)
        self.assertIn('&lt;script&gt;',page)
        self.assertIn('未分类',page)
        with self.assertRaises(ValueError):render(data,[{'unit_id':'fake','intent':'L1','rationale':'x'}])
        note={'unit_id':data['units'][0]['unit_id'],'intent':'L3','rationale':'Brand discovery'}
        self.assertIn('Brand discovery',render(data,[note]))
        with self.assertRaises(ValueError):render(data,[note,note])


if __name__=='__main__':
    unittest.main()
